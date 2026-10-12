import openpyxl, json, csv, sys
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter as L
import os; U=os.environ.get('DATA_DIR','./')
SRC=U+'Electric_Cars_Purchase_Intention-Data_in_Brief.csv'
ref=json.load(open('tb_ref.json'))
F=lambda **k: Font(**{'name':'Arial','size':10,**k})
HB=PatternFill('solid',fgColor='D9E1F2'); YL=PatternFill('solid',fgColor='FFF2CC')
thin=Side(style='thin',color='BFBFBF'); BOX=Border(left=thin,right=thin,top=thin,bottom=thin)
def hdr(ws,row,vals,col=1):
    for j,v in enumerate(vals):
        c=ws.cell(row,col+j,v); c.font=F(bold=True); c.fill=HB; c.border=BOX
        c.alignment=Alignment(wrap_text=True,vertical='center',horizontal='center')
def put(ws,r,c,v,fmt=None,bold=False):
    x=ws.cell(r,c,v); x.font=F(bold=bold); x.border=BOX
    if fmt: x.number_format=fmt
    return x
wb=openpyxl.Workbook(); wb.remove(wb.active)
ITEMS=['TB1','TB2','TB3','TB4','TB5','TB6']
TXT={'TB1':'I think e-cars are sleek','TB2':'I think e-cars are durable','TB3':'I think e-cars are available in different size, colours and designs',
     'TB4':'I think e-cars are aesthetically appealing','TB5':'I think e-cars are technologically superior','TB6':'I think e-cars give comfortable ride'}

# ---------- README ----------
rd=wb.create_sheet('README')
lines=[("Soal 7 - Item analysis and reliability of the TB scale (TB1-TB6)",True),
("Data source (verified from the uploaded files only)",True),
("File 1: Electric_Cars_Purchase_Intention-Data_in_Brief.csv (MD5 09e51b6cd9256e2b3c7d0faef39daf9d), 510 data rows x 41 columns.",False),
("File 2: Electric_Cars_Purchase_Intention-Maheswaran_and_Veena.sav (MD5 5bbeb8c35e913cba2aa6b264b5a7a33b), SPSS 23, created 21 May 2021; item values identical to File 1.",False),
("Dataset landing page named by the user: https://data.mendeley.com/datasets/tz7krp8rvz/2 (could not be opened from the analysis environment; licence, version and citation NOT verified).",False),
("Item wording below comes from the SPSS variable labels. The SPSS file has NO value labels for the Likert items: the meaning of codes 1-5 is NOT documented in the files.",False),
("Construct name for 'TB' is NOT documented in the files.",False),
("Sheets",True),
("Raw_Data: the CSV exactly as uploaded (blank cells kept blank). Do not edit.",False),
("Data_Check: range, missing-value and duplicate checks computed from Raw_Data.",False),
("TB_Primary: TB1-TB6 linked from Raw_Data (n = 510) with total, rest scores and odd/even half scores.",False),
("TB_Sensitivity: same, excluding Raw_Data row 142 (data row 141), the second of two rows identical on all 41 variables (n = 509).",False),
("Item_Analysis: corrected item-total correlations, alpha if item deleted, inter-item correlations.",False),
("Reliability: Cronbach's alpha (with Feldt 95% CI) and odd-even split-half with Spearman-Brown correction.",False),
("Cross_Check: Excel results vs an independent Python (pandas/scipy) computation.",False),
("Pre-specified decisions (fixed before computing)",True),
("Items: TB1-TB6 chosen on content (all six rate attributes of the e-car itself). No item is dropped after seeing results.",False),
("Item order 1-6 = TB1..TB6 (SPSS variable order). Odd half = TB1+TB3+TB5; even half = TB2+TB4+TB6.",False),
("Primary analysis n = 510 (as published); sensitivity analysis n = 509. The identical rows are NOT asserted to be an entry error.",False),
("Criteria (assumed; official Soal 7 criteria not available): corrected item-total r >= 0.30; alpha / Spearman-Brown >= 0.70 acceptable, 0.60-0.69 questionable, < 0.60 poor.",False)]
for i,(t,b) in enumerate(lines,1):
    c=rd.cell(i,1,t); c.font=F(bold=b,size=12 if i==1 else 10)
rd.column_dimensions['A'].width=150

# ---------- Raw_Data ----------
ws=wb.create_sheet('Raw_Data')
with open(SRC,encoding='utf-8-sig',newline='') as f:
    rows=list(csv.reader(f))
hdrs=rows[0]
for j,h in enumerate(hdrs,1):
    c=ws.cell(1,j,h); c.font=F(bold=True); c.fill=HB
for i,row in enumerate(rows[1:],2):
    for j,v in enumerate(row,1):
        v=v.strip()
        if v=='': continue
        ws.cell(i,j,int(v) if v.lstrip('-').isdigit() else float(v)).font=F()
ws.freeze_panes='A2'
NR=len(rows)-1; LAST=NR+1   # 511
col={h:L(j) for j,h in enumerate(hdrs,1)}

# ---------- Data_Check ----------
dc=wb.create_sheet('Data_Check')
dc['A1']='Data checks (computed from Raw_Data)'; dc['A1'].font=F(bold=True,size=12)
hdr(dc,3,['Variable','Item wording','N non-blank','Blank','Min','Max','Outside 1-5','n=1','n=2','n=3','n=4','n=5','Mean','SD'])
r=4
for it in ITEMS:
    g=f"Raw_Data!${col[it]}$2:${col[it]}${LAST}"
    vals=[it,TXT[it],f'=COUNT({g})',f'=COUNTBLANK({g})',f'=MIN({g})',f'=MAX({g})',f'=COUNT({g})-COUNTIFS({g},">=1",{g},"<=5")']+[f'=COUNTIF({g},{k})' for k in range(1,6)]+[f'=AVERAGE({g})',f'=STDEV({g})']
    for j,v in enumerate(vals,1): put(dc,r,j,v,'0.000' if j>=13 else None)
    r+=1
r+=1
dc.cell(r,1,'All 31 Likert items (CC1:TB6)').font=F(bold=True); r+=1
allg=f"Raw_Data!$K$2:$AO${LAST}"
for lab,f in [('Cells',f'=ROWS({allg})*COLUMNS({allg})'),('Blank cells',f'=COUNTBLANK({allg})'),('Cells outside 1-5',f'=COUNT({allg})-COUNTIFS({allg},">=1",{allg},"<=5")'),('Data rows',f'=COUNTA(Raw_Data!$J$2:$J${LAST})')]:
    dc.cell(r,1,lab).font=F(); dc.cell(r,3,f).font=F(); r+=1
r+=1
dc.cell(r,1,'Duplicate rows (all 41 variables identical)').font=F(bold=True); r+=1
dc.cell(r,1,'Rows that have an identical twin').font=F(); dc.cell(r,3,f'=COUNTIF($Q$4:$Q${LAST+2},">1")').font=F(); r+=1
dc.cell(r,1,'Raw_Data rows involved').font=F(); dc.cell(r,3,'142 and 139 (data rows 141 and 138); see column Q').font=F(); r+=1
dc.cell(r,1,'Age outside 18-35 (count)').font=F(); dc.cell(r,3,f'=COUNTIF(Raw_Data!$J$2:$J${LAST},"<18")+COUNTIF(Raw_Data!$J$2:$J${LAST},">35")').font=F(); r+=1
dc.cell(r,1,'Note: Age and the demographic coding issues do not enter the TB analysis.').font=F(italic=True)
# helper keys in P/Q aligned with Raw_Data rows (row 4 here = Raw_Data row 2)
dc['P3']='Row key (Raw_Data row)'; dc['Q3']='Count of identical rows'; dc['O3']='Raw_Data row'
for c_ in ('O3','P3','Q3'): dc[c_].font=F(bold=True); dc[c_].fill=HB
for i in range(2,LAST+1):
    rr=i+2
    dc.cell(rr,15,i).font=F()
    key='&"|"&'.join(f'Raw_Data!{L(j)}{i}' for j in range(1,42))
    dc.cell(rr,16,'='+key).font=F()
    dc.cell(rr,17,f'=COUNTIF($P$4:$P${LAST+2},P{rr})').font=F()
dc.column_dimensions['A'].width=38; dc.column_dimensions['B'].width=52; dc.column_dimensions['P'].width=30
for j in range(3,15): dc.column_dimensions[L(j)].width=9

# ---------- TB data sheets ----------
def tb_sheet(name, rawrows, note):
    t=wb.create_sheet(name)
    H=['Raw_Data row']+ITEMS+['Total']+[f'Rest_{x}' for x in ITEMS]+['Odd (TB1+TB3+TB5)','Even (TB2+TB4+TB6)']
    hdr(t,1,H)
    for k,i in enumerate(rawrows):
        r=k+2
        t.cell(r,1,i)
        for j,it in enumerate(ITEMS):
            t.cell(r,2+j,f'=Raw_Data!{col[it]}{i}')
        t.cell(r,8,f'=SUM(B{r}:G{r})')
        for j in range(6):
            t.cell(r,9+j,f'=H{r}-{L(2+j)}{r}')
        t.cell(r,15,f'=B{r}+D{r}+F{r}'); t.cell(r,16,f'=C{r}+E{r}+G{r}')
    last=len(rawrows)+1
    t.cell(last+2,1,note).font=F(italic=True)
    t.freeze_panes='B2'
    for j in range(1,17): t.column_dimensions[L(j)].width=11
    return last
LP=tb_sheet('TB_Primary',list(range(2,LAST+1)),'All 510 data rows (Raw_Data rows 2-511).')
LS=tb_sheet('TB_Sensitivity',[i for i in range(2,LAST+1) if i!=142],'Raw_Data row 142 (data row 141) excluded: it is identical on all 41 variables to row 139 (data row 138). This is a sensitivity check, not a claim that the row is an error.')

# ---------- Item_Analysis ----------
ia=wb.create_sheet('Item_Analysis')
ia['A1']='Item analysis: corrected item-total correlation (CITC) = Pearson r between an item and the sum of the other five TB items'; ia['A1'].font=F(bold=True,size=12)
ia['A2']='CITC criterion (assumed; edit if the lecturer specifies another)'; ia['F2']=0.30; ia['F2'].font=F(color='0000FF'); ia['F2'].fill=YL
ia['A3']='Alpha for significance tests'; ia['F3']=0.05; ia['F3'].font=F(color='0000FF'); ia['F3'].fill=YL
def item_block(r0, sheet, last, title):
    ia.cell(r0,1,title).font=F(bold=True)
    ia.cell(r0+1,1,'n'); ia.cell(r0+1,2,f'=COUNT({sheet}!$H$2:$H${last})')
    ia.cell(r0+2,1,'r critical (two-tailed)'); ia.cell(r0+2,2,f'=TINV($F$3,B{r0+1}-2)/SQRT(TINV($F$3,B{r0+1}-2)^2+B{r0+1}-2)'); ia.cell(r0+2,2).number_format='0.000'
    h=r0+4
    hdr(ia,h,['Item','Wording','Mean','SD','Variance','r item-total (uncorrected)','CITC','t','p (two-tailed)','Alpha if item deleted','Meets CITC >= criterion?','Statistically significant?'])
    for j,it in enumerate(ITEMS):
        r=h+1+j; ic=L(2+j); rc=L(9+j)
        g=f'{sheet}!${ic}$2:${ic}${last}'; rest=f'{sheet}!${rc}$2:${rc}${last}'; tot=f'{sheet}!$H$2:$H${last}'
        vals=[it,TXT[it],f'=AVERAGE({g})',f'=STDEV({g})',f'=VAR({g})',f'=CORREL({g},{tot})',f'=CORREL({g},{rest})',
              f'=G{r}*SQRT(($B${r0+1}-2)/(1-G{r}^2))',f'=TDIST(ABS(H{r}),$B${r0+1}-2,2)',
              # alpha if item deleted: k=5, sum var of others = SUM(E)-E_r, variance of rest score
              f'=(5/4)*(1-(SUM($E${h+1}:$E${h+6})-E{r})/VAR({rest}))',
              f'=IF(G{r}>=$F$2,"Yes","No")',f'=IF(I{r}<$F$3,"Yes","No")']
        for k,v in enumerate(vals,1):
            put(ia,r,k,v,{3:'0.000',4:'0.000',5:'0.000',6:'0.000',7:'0.000',8:'0.00',9:'0.0000',10:'0.000'}.get(k))
    s=h+7
    ia.cell(s,1,'Items meeting CITC criterion'); ia.cell(s,3,f'=COUNTIF(K{h+1}:K{h+6},"Yes")&" of 6"')
    # inter-item matrix
    m=s+2; ia.cell(m,1,'Inter-item Pearson correlations').font=F(bold=True)
    hdr(ia,m+1,['']+ITEMS)
    for a,ia_ in enumerate(ITEMS):
        ia.cell(m+2+a,1,ia_).font=F(bold=True)
        for b,ib in enumerate(ITEMS):
            ga=f'{sheet}!${L(2+a)}$2:${L(2+a)}${last}'; gb=f'{sheet}!${L(2+b)}$2:${L(2+b)}${last}'
            put(ia,m+2+a,2+b,f'=CORREL({ga},{gb})','0.000')
    ia.cell(m+9,1,'Mean inter-item r (15 pairs)'); ia.cell(m+9,3,f'=(SUM(B{m+2}:G{m+7})-6)/30'); ia.cell(m+9,3).number_format='0.000'
    ia.cell(m+10,1,'Smallest inter-item r'); ia.cell(m+10,3,f'=MIN(B{m+2}:G{m+7})'); ia.cell(m+10,3).number_format='0.000'
    return h, m, m+11
P_h,P_m,nxt=item_block(5,'TB_Primary',LP,'PRIMARY ANALYSIS (n = 510, as published)')
S_h,S_m,_=item_block(nxt+2,'TB_Sensitivity',LS,'SENSITIVITY ANALYSIS (n = 509, second identical row removed)')
for j,w in enumerate([10,58,8,8,9,13,9,8,11,12,12,12],1): ia.column_dimensions[L(j)].width=w

# ---------- Reliability ----------
rl=wb.create_sheet('Reliability')
rl['A1']="Reliability of the 6-item TB scale"; rl['A1'].font=F(bold=True,size=12)
rl['A2']='CI level'; rl['C2']=0.95; rl['C2'].font=F(color='0000FF'); rl['C2'].fill=YL
rl['A3']='Acceptable threshold (assumed)'; rl['C3']=0.70; rl['C3'].font=F(color='0000FF'); rl['C3'].fill=YL
rl['A4']='Questionable threshold (assumed)'; rl['C4']=0.60; rl['C4'].font=F(color='0000FF'); rl['C4'].fill=YL
hdr(rl,6,['Statistic','Primary (n = 510)','Sensitivity (n = 509)','Formula / note'])
def R(sheet,last,iah):
    T=f'{sheet}!$H$2:$H${last}'; O=f'{sheet}!$O$2:$O${last}'; E=f'{sheet}!$P$2:$P${last}'
    return [f'=COUNT({T})',6,f'=SUM(Item_Analysis!$E${iah+1}:$E${iah+6})',f'=VAR({T})',
            None,None,None,None,  # alpha rows filled below with relative refs
            f'=CORREL({O},{E})',None,None,f'=VAR({O})',f'=VAR({E})',None]
labels=['n','k (items)','Sum of item variances','Variance of total score',"Cronbach's alpha",'Alpha 95% CI lower (Feldt)','Alpha 95% CI upper (Feldt)','Alpha decision',
        'r between odd and even half scores','Spearman-Brown corrected split-half','Split-half decision','Variance of odd half','Variance of even half','Guttman split-half (lambda-4, same split; supplementary)']
notes=['COUNT of total scores','TB1-TB6','SUM of VAR of each item (Item_Analysis col E)','VAR of Total',
       'k/(k-1) * (1 - sum item var / total var)','1-(1-alpha)*FINV(0.025, n-1, (n-1)(k-1))','1-(1-alpha)*FINV(0.975, n-1, (n-1)(k-1))','>= 0.70 acceptable; 0.60-0.69 questionable; < 0.60 poor',
       'Odd = TB1+TB3+TB5, Even = TB2+TB4+TB6','2r / (1 + r)','same thresholds','','','2*(1-(var odd + var even)/var total)']
for i,lab in enumerate(labels):
    r=7+i; rl.cell(r,1,lab).font=F(bold=i in (4,9)); rl.cell(r,4,notes[i]).font=F(italic=True)
for cc,(sheet,last,iah) in zip(('B','C'),(('TB_Primary',LP,P_h),('TB_Sensitivity',LS,S_h))):
    base=R(sheet,last,iah)
    f={7:base[0],8:base[1],9:base[2],10:base[3],
       11:f'=({cc}8/({cc}8-1))*(1-{cc}9/{cc}10)',
       12:f'=1-(1-{cc}11)*FINV((1-$C$2)/2,{cc}7-1,({cc}7-1)*({cc}8-1))',
       13:f'=1-(1-{cc}11)*FINV(1-(1-$C$2)/2,{cc}7-1,({cc}7-1)*({cc}8-1))',
       14:f'=IF({cc}11>=$C$3,"Acceptable",IF({cc}11>=$C$4,"Questionable","Poor"))',
       15:base[8],16:f'=2*{cc}15/(1+{cc}15)',
       17:f'=IF({cc}16>=$C$3,"Acceptable",IF({cc}16>=$C$4,"Questionable","Poor"))',
       18:base[11],19:base[12],20:f'=2*(1-({cc}18+{cc}19)/{cc}10)'}
    for r,v in f.items():
        put(rl,r,{'B':2,'C':3}[cc],v,None if r in (7,8,14,17) else '0.000',bold=r in (11,16))
rl.column_dimensions['A'].width=48; rl.column_dimensions['B'].width=18; rl.column_dimensions['C'].width=20; rl.column_dimensions['D'].width=60

# ---------- Cross_Check ----------
cx=wb.create_sheet('Cross_Check')
cx['A1']='Cross-check: Excel formulas vs independent Python computation (pandas 2.x / scipy). Python values are pasted as fixed numbers.'; cx['A1'].font=F(bold=True)
hdr(cx,3,['Analysis','Statistic','Excel (formula)','Python (fixed value)','Absolute difference'])
r=4
for o,label,h,cc in ((ref[0],'Primary n=510',P_h,'B'),(ref[1],'Sensitivity n=509',S_h,'C')):
    for j,it in enumerate(o['items']):
        for stat,colx in (('citc','G'),('alpha_if_deleted','J'),('mean','C'),('sd','D')):
            put(cx,r,1,label); put(cx,r,2,f"{it['item']} {stat}"); put(cx,r,3,f'=Item_Analysis!{colx}{h+1+j}','0.000000')
            put(cx,r,4,it[stat],'0.000000'); put(cx,r,5,f'=ABS(C{r}-D{r})','0.0E+00'); r+=1
    for stat,cell in (('alpha',11),('spearman_brown',16),('r_half',15),('guttman',20)):
        put(cx,r,1,label); put(cx,r,2,stat); put(cx,r,3,f'=Reliability!{cc}{cell}','0.000000'); put(cx,r,4,o[stat],'0.000000'); put(cx,r,5,f'=ABS(C{r}-D{r})','0.0E+00'); r+=1
    for stat,cell,idx in (('alpha CI lower',12,0),('alpha CI upper',13,1)):
        put(cx,r,1,label); put(cx,r,2,stat); put(cx,r,3,f'=Reliability!{cc}{cell}','0.000000'); put(cx,r,4,o['alpha_ci'][idx],'0.000000'); put(cx,r,5,f'=ABS(C{r}-D{r})','0.0E+00'); r+=1
cx.cell(r+1,1,'Largest absolute difference'); cx.cell(r+1,5,f'=MAX(E4:E{r-1})').number_format='0.0E+00'
cx.cell(r+2,1,'Excel values in this file were computed by LibreOffice (headless) and have not been opened in Microsoft Excel.').font=F(italic=True)
for j,w in enumerate([20,26,18,20,18],1): cx.column_dimensions[L(j)].width=w

wb._sheets=[wb[n] for n in ['README','Item_Analysis','Reliability','Cross_Check','Data_Check','TB_Primary','TB_Sensitivity','Raw_Data']]
wb.active=0
wb.save(sys.argv[1]); print('saved',P_h,S_h)
