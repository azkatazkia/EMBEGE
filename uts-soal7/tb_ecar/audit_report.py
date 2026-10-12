"""Check that every key number quoted in the report matches the workbook (LibreOffice-computed values)."""
import openpyxl, sys, re
wb = openpyxl.load_workbook(sys.argv[1], data_only=True)
doc = open(sys.argv[2], encoding='utf-8').read()
ia, rl = wb['Item_Analysis'], wb['Reliability']
fails = []
def chk(label, value, text, dp):
    s = f"{value:.{dp}f}"
    ok = s in doc if text is None else (text in doc and s in text)
    print(f"{'OK ' if ok else 'FAIL'} {label:42s} workbook={value:.6f} expected in report as '{s}'")
    if not ok: fails.append(label)
P = 10; S = 36  # first item row of primary / sensitivity blocks
for j, it in enumerate(['TB1','TB2','TB3','TB4','TB5','TB6']):
    r = P + j
    row = re.search(rf"\| {it} [^\n]*", doc).group(0)
    for col, name, dp in ((3,'mean',3),(4,'SD',3),(6,'r uncorrected',3),(7,'CITC',3),(10,'alpha if deleted',3)):
        chk(f"{it} {name}", ia.cell(r, col).value, row, dp)
    pv = ia.cell(r, 9).value
    exp = '< 0.001' if pv < 0.001 else f"{pv:.3f}"
    ok = exp in row; print(f"{'OK ' if ok else 'FAIL'} {it} p (CITC){'':31s} workbook={pv:.2e} expected '{exp}'")
    if not ok: fails.append(it+' p')
    meets = 'Yes' if ia.cell(r, 7).value >= 0.30 else 'No'
    ok = row.rstrip().endswith(f"| {meets} |"); print(f"{'OK ' if ok else 'FAIL'} {it} CITC>=0.30 flag = {meets}")
    if not ok: fails.append(it+' flag')
chk('alpha (n=510)', rl['B11'].value, None, 3)
chk('alpha CI lower', rl['B12'].value, None, 3)
chk('alpha CI upper', rl['B13'].value, None, 3)
chk('r_hh', rl['B15'].value, None, 3)
chk('Spearman-Brown', rl['B16'].value, None, 3)
chk('Guttman', rl['B20'].value, None, 3)
chk('alpha 4dp n=510', rl['B11'].value, None, 4)
chk('alpha 4dp n=509', rl['C11'].value, None, 4)
chk('SB 4dp n=510', rl['B16'].value, None, 4)
chk('SB 4dp n=509', rl['C16'].value, None, 4)
chk('var odd', rl['B18'].value, None, 2)
chk('var even', rl['B19'].value, None, 2)
mr = ia['C27'].value if ia['A27'].value and 'Mean' in str(ia['A27'].value) else None
# locate mean inter-item r cell
for r in range(1, 40):
    if str(ia.cell(r,1).value).startswith('Mean inter-item'): chk('mean inter-item r', ia.cell(r,3).value, None, 3); break
cit = [ia.cell(S+j,7).value for j in range(6)]
chk('n=509 CITC min', min(cit), None, 3); chk('n=509 CITC max', max(cit), None, 3)
d = max(abs(ia.cell(P+j,c).value - ia.cell(S+j,c).value) for j in range(6) for c in (3,4,6,7,10))
print(f"{'OK ' if d <= 0.004 else 'FAIL'} max |n510-n509| difference across item stats = {d:.4f} (report says <= 0.004)")
if d > 0.004: fails.append('sensitivity bound')
print('\nFAILURES:', fails if fails else 'none')
